"""Render a publication-quality Word results table from the per-phenotype receipts.

Reads paper/results/<phenotype>_build_receipt.json (written by aggregate_results.py)
and builds a formatted .docx table — bordered grid, shaded bold header row, numbered
caption, and a methods/abbreviations footnote — suitable to send to a collaborator.

Usage:
    python pipeline/make_results_table_docx.py                       # -> ~/Downloads/results table.docx
    python pipeline/make_results_table_docx.py --out "/path/x.docx"
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = REPO_ROOT / "paper" / "results"

# Display order + pretty names (clinical grouping: cardiovascular, then metabolic/renal, then other).
ORDER = [
    ("stroke", "Ischemic stroke"),
    ("ischemic_heart_disease", "Ischemic heart disease"),
    ("ascvd", "Atherosclerotic CVD (ASCVD)"),
    ("heart_failure", "Heart failure"),
    ("hypertension", "Hypertension"),
    ("dyslipidemia", "Dyslipidemia"),
    ("type2_diabetes", "Type 2 diabetes"),
    ("prediabetes", "Prediabetes"),
    ("chronic_kidney_disease", "Chronic kidney disease"),
    ("obstructive_sleep_apnea", "Obstructive sleep apnea"),
    ("asthma", "Asthma"),
]
COLS = ["Phenotype", "Cases /\nControls", "Features", "ROC AUC", "Accuracy", "Precision", "Recall", "F1"]


def _shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _set_cell(cell, text, *, bold=False, align="center", size=9, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    for i, line in enumerate(str(text).split("\n")):
        run = p.add_run(("\n" if i else "") + line)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        if color:
            run.font.color.rgb = RGBColor(*color)


def _fmt(v, nd=3):
    if v is None:
        return "—"
    return f"{v:.{nd}f}" if isinstance(v, float) else str(v)


def load_rows():
    rows = []
    for key, pretty in ORDER:
        f = RESULTS_DIR / f"{key}_build_receipt.json"
        if not f.exists():
            continue
        r = json.loads(f.read_text())
        if not r.get("has_serialized_model"):
            continue
        rows.append({
            "name": pretty,
            "cc": f"{r.get('n_positive','—')} / {r.get('n_negative','—')}",
            "feat": r.get("n_features_selected", "—"),
            "auc": r.get("roc_auc"), "acc": r.get("accuracy"),
            "prec": r.get("precision"), "rec": r.get("recall"), "f1": r.get("f1"),
        })
    return rows


def build(out_path: Path):
    rows = load_rows()
    if not rows:
        raise SystemExit("No build receipts with a model found in paper/results/. Run a build first.")

    doc = Document()
    # page + base style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # Caption
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(6)
    r1 = cap.add_run("Table 1. ")
    r1.bold = True
    r1.font.size = Pt(10)
    r2 = cap.add_run(
        "Performance of phenotype classification models built through the i2b2-ML JSON API "
        "on MIMIC-IV. Each model was trained and scored entirely by the i2b2-ML engine "
        "(POST /etl/job, jobType:ml); values are the engine's held-out test-set metrics."
    )
    r2.font.size = Pt(10)

    # Table
    table = doc.add_table(rows=1, cols=len(COLS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # header
    hdr = table.rows[0].cells
    for i, c in enumerate(COLS):
        _set_cell(hdr[i], c, bold=True, align=("left" if i == 0 else "center"), size=9)
        _shade(hdr[i], "1F3864")  # dark blue
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # data
    for idx, r in enumerate(rows):
        cells = table.add_row().cells
        vals = [r["name"], r["cc"], r["feat"], _fmt(r["auc"]), _fmt(r["acc"]),
                _fmt(r["prec"]), _fmt(r["rec"]), _fmt(r["f1"])]
        for i, v in enumerate(vals):
            _set_cell(cells[i], v, align=("left" if i == 0 else "center"), size=9,
                      bold=(i == 3))  # bold the headline ROC AUC column
        if idx % 2 == 1:
            for c in cells:
                _shade(c, "EAF0F8")  # light blue zebra

    # mean AUROC summary row
    aucs = [r["auc"] for r in rows if isinstance(r["auc"], float)]
    if aucs:
        cells = table.add_row().cells
        _set_cell(cells[0], f"Mean (n={len(rows)} phenotypes)", bold=True, align="left", size=9)
        for i in range(1, 3):
            _set_cell(cells[i], "", size=9)
        _set_cell(cells[3], _fmt(sum(aucs) / len(aucs)), bold=True, size=9)
        for i in range(4, 8):
            _set_cell(cells[i], "", size=9)
        for c in cells:
            _shade(c, "D9D9D9")

    # footnote
    fn = doc.add_paragraph()
    fn.paragraph_format.space_before = Pt(6)
    note = fn.add_run(
        "Models: elastic-net logistic regression with SMOTE oversampling and 5-fold "
        "cross-validated hyperparameter selection, on a 50% held-out test split. Cases are "
        "ICD-defined (silver standard); 32 structured features (demographics, comorbidities, "
        "laboratory values, vitals, medication classes) were supplied, of which the engine "
        "retained those shown after dropping high-missingness variables. The comorbidity feature "
        "matching each outcome was excluded to avoid label leakage. ROC AUC, area under the "
        "receiver-operating-characteristic curve."
    )
    note.font.size = Pt(8)
    note.italic = True
    note.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}  ({len(rows)} phenotypes)")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(Path.home() / "Downloads" / "results table.docx"))
    args = ap.parse_args(argv)
    build(Path(args.out))
    # also keep a copy in the repo
    build(REPO_ROOT / "paper" / "tables" / "results_table.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
